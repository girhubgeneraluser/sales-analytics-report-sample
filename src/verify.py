"""Independent QA for the Northwind performance report.

Sample project - synthetic data, not client work.

Recomputes the headline KPIs straight from data/raw/ using ONLY the Python standard library
(csv, datetime, zoneinfo) - no pandas, no shared code with clean.py / analyze.py - and checks
them against the numbers written into the Excel workbook and the executive summary.

Exit code 0 = all checks passed, 1 = at least one failure.
"""
from __future__ import annotations

import csv
import os
import datetime as dt
import re
import sys
from collections import defaultdict
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook  # only used to READ the deliverable

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
XLSX = Path(os.environ.get("VERIFY_XLSX", ROOT / "deliverables" / "Northwind_Performance_Report.xlsx"))
SUMMARY_MD = ROOT / "deliverables" / "executive_summary.md"
NY = ZoneInfo("America/New_York")
UTC = dt.timezone.utc

results: list[tuple[bool, str]] = []


def check(ok: bool, msg: str) -> None:
    results.append((bool(ok), msg))


def close(a: float, b: float, rel: float = 1e-6, abs_: float = 0.01) -> bool:
    return abs(a - b) <= max(abs_, rel * max(abs(a), abs(b)))


# =========================================================================== recompute from raw
def parse_ts(s: str) -> dt.datetime:
    s = s.strip()
    if s.endswith("Z"):
        return dt.datetime.strptime(s, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=UTC)
    if re.search(r"[+-]\d\d:\d\d$", s):
        return dt.datetime.fromisoformat(s).astimezone(UTC)
    return dt.datetime.strptime(s, "%Y-%m-%d %H:%M:%S").replace(tzinfo=NY).astimezone(UTC)


def recompute() -> dict:
    fx = {r["month"]: float(r["cad_to_usd"]) for r in csv.DictReader(open(RAW / "fx_rates_monthly.csv"))}
    cost = {r["sku"]: float(r["unit_cost_usd"]) for r in csv.DictReader(open(RAW / "products.csv"))}
    cat = {r["sku"]: r["category"] for r in csv.DictReader(open(RAW / "products.csv"))}

    seen = set()
    kept = []
    for r in csv.DictReader(open(RAW / "shopify_orders_export.csv")):
        key = (r["order_id"], r["line_number"])
        if key in seen:
            continue
        seen.add(key)
        email = r["customer_email"].strip().lower()
        if "@northwind-test." in email:
            continue
        if r["financial_status"] in ("voided", "cancelled"):
            continue
        if int(r["quantity"]) <= 0:
            continue
        r["email"] = email
        kept.append(r)

    # recover missing emails from sibling lines
    order_email = {}
    for r in kept:
        if r["email"] and r["order_id"] not in order_email:
            order_email[r["order_id"]] = r["email"]
    for r in kept:
        if not r["email"]:
            r["email"] = order_email.get(r["order_id"], "")

    orders: dict[str, dict] = {}
    product_net = defaultdict(float)
    lines_by_order_cat = defaultdict(set)
    for r in kept:
        ts = parse_ts(r["created_at"])
        month = ts.astimezone(NY).strftime("%Y-%m")
        rate = fx[month] if r["currency"].strip().upper() == "CAD" else 1.0
        q = int(r["quantity"])
        gross = q * float(r["unit_price"]) * rate
        disc = float(r["line_discount"]) * rate
        ref = float(r["refunded_amount"]) * rate
        o = orders.setdefault(r["order_id"], dict(ts=ts, month=month, email=r["email"], channel=r["source_channel"],
                                                  gross=0.0, disc=0.0, ref=0.0, cogs=0.0))
        o["gross"] += gross
        o["disc"] += disc
        o["ref"] += ref
        o["cogs"] += q * cost[r["sku"]]
        product_net[r["sku"]] += gross - disc - ref
        lines_by_order_cat[r["order_id"]].add(cat[r["sku"]])

    spend_seen = set()
    spend_by_ch = defaultdict(float)
    spend_by_month = defaultdict(float)
    spend_by_ch_month = defaultdict(float)
    for r in csv.DictReader(open(RAW / "ad_spend_daily.csv")):
        key = (r["date"], r["channel"], r["campaign"])
        if key in spend_seen:
            continue
        spend_seen.add(key)
        v = float(r["spend_usd"])
        spend_by_ch[r["channel"]] += v
        spend_by_month[r["date"][:7]] += v
        spend_by_ch_month[(r["channel"], r["date"][:7])] += v

    # customers: first order = earliest timestamp (tie -> lowest order id)
    by_cust = defaultdict(list)
    for oid, o in orders.items():
        if o["email"]:
            by_cust[o["email"]].append((o["ts"], int(oid), oid))
    first_order = {}
    first_month = {}
    for e, lst in by_cust.items():
        lst.sort()
        first_order[e] = lst[0][2]
        first_month[e] = orders[lst[0][2]]["month"]
    new_by_ch = defaultdict(int)
    new_by_ch_month = defaultdict(int)
    for e, oid in first_order.items():
        new_by_ch[orders[oid]["channel"]] += 1
        new_by_ch_month[(orders[oid]["channel"], orders[oid]["month"])] += 1

    net = {oid: o["gross"] - o["disc"] - o["ref"] for oid, o in orders.items()}
    rev_by_ch = defaultdict(float)
    rev_by_ch_month = defaultdict(float)
    m_net = defaultdict(float)
    m_orders = defaultdict(int)
    m_cust = defaultdict(set)
    m_ret = defaultdict(set)
    for oid, o in orders.items():
        rev_by_ch[o["channel"]] += net[oid]
        rev_by_ch_month[(o["channel"], o["month"])] += net[oid]
        m_net[o["month"]] += net[oid]
        m_orders[o["month"]] += 1
        if o["email"]:
            m_cust[o["month"]].add(o["email"])
            if first_month[o["email"]] < o["month"]:
                m_ret[o["month"]].add(o["email"])

    # cohorts
    def midx(m: str) -> int:
        y, mo = m.split("-")
        return int(y) * 12 + int(mo)

    cohort_size = defaultdict(int)
    for e, m in first_month.items():
        cohort_size[m] += 1
    cohort_active = defaultdict(set)
    for oid, o in orders.items():
        if o["email"]:
            cohort_active[(first_month[o["email"]], midx(o["month"]) - midx(first_month[o["email"]]))].add(o["email"])

    # bedding-first repeat rates
    n_orders = {e: len(v) for e, v in by_cust.items()}
    bed = [n_orders[e] >= 2 for e, oid in first_order.items() if "Bedding" in lines_by_order_cat[oid]]
    oth = [n_orders[e] >= 2 for e, oid in first_order.items() if "Bedding" not in lines_by_order_cat[oid]]

    tot_net = sum(net.values())
    tot_gross = sum(o["gross"] for o in orders.values())
    tot_disc = sum(o["disc"] for o in orders.values())
    tot_ref = sum(o["ref"] for o in orders.values())
    tot_cogs = sum(o["cogs"] for o in orders.values())
    tot_spend = sum(spend_by_ch.values())
    k = {
        "Net revenue": tot_net, "Gross sales": tot_gross, "Discounts": tot_disc, "Refunds": tot_ref,
        "Orders": len(orders), "AOV": tot_net / len(orders), "Identified customers": len(by_cust),
        "Repeat-customer rate": sum(1 for v in by_cust.values() if len(v) >= 2) / len(by_cust),
        "New customers": len(first_order), "Ad spend (Meta + Google)": tot_spend,
        "MER (net revenue ÷ ad spend)": tot_net / tot_spend, "Blended CAC": tot_spend / len(first_order),
        "Meta ROAS": rev_by_ch["meta"] / spend_by_ch["meta"], "Google ROAS": rev_by_ch["google"] / spend_by_ch["google"],
        "Meta CAC": spend_by_ch["meta"] / new_by_ch["meta"], "Google CAC": spend_by_ch["google"] / new_by_ch["google"],
        "Refund rate": tot_ref / (tot_gross - tot_disc), "Gross margin %": (tot_net - tot_cogs) / tot_net,
    }
    months = sorted(m_net)
    return dict(k=k, months=months, m_net=m_net, m_orders=m_orders, m_cust=m_cust, m_ret=m_ret,
                cohort_size=cohort_size, cohort_active=cohort_active, product_net=product_net,
                rev_by_ch=rev_by_ch, spend_by_ch=spend_by_ch, new_by_ch=new_by_ch, spend_by_month=spend_by_month,
                rev_by_ch_month=rev_by_ch_month, spend_by_ch_month=spend_by_ch_month,
                new_by_ch_month=new_by_ch_month, rep_bed=sum(bed) / len(bed), rep_oth=sum(oth) / len(oth),
                n_clean_lines=len(kept))


# =========================================================================== read workbook
def table(ws, header_row: int) -> list[dict]:
    hdr = [c.value for c in ws[header_row]]
    out = []
    r = header_row + 1
    while ws.cell(row=r, column=1).value not in (None, "Total"):
        out.append({h: ws.cell(row=r, column=j + 1).value for j, h in enumerate(hdr) if h})
        r += 1
    return out


def find_row(ws, text: str, col: int = 1) -> int:
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=col).value == text:
            return r
    raise KeyError(text)


def main() -> int:
    R = recompute()
    k = R["k"]
    wb = load_workbook(XLSX)  # values (formulas only on Total rows, not read here)

    # ---------------- 1. headline KPIs vs Summary sheet
    ws = wb["Summary KPIs"]
    wk = {}
    for r in range(5, 5 + len(k)):
        wk[ws.cell(row=r, column=1).value] = ws.cell(row=r, column=2).value
    for name, val in k.items():
        got = wk.get(name)
        check(got is not None and close(float(got), float(val)),
              f"Summary KPI '{name}': workbook={float(got):,.4f}  recomputed={val:,.4f}" if got is not None
              else f"Summary KPI '{name}' missing from workbook")

    # ---------------- 2. monthly
    ws = wb["Monthly"]
    mrows = table(ws, 4)
    check(len(mrows) == 18, f"Monthly sheet has 18 months (found {len(mrows)})")
    for row in mrows:
        m = row["Month"]
        check(close(row["Net Revenue"], R["m_net"][m]) and row["Orders"] == R["m_orders"][m],
              f"Monthly {m}: net revenue & orders match raw recompute")
        rr = len(R["m_ret"][m]) / len(R["m_cust"][m])
        check(close(row["Repeat Customer Rate"], rr), f"Monthly {m}: repeat-customer rate {rr:.4f} matches")
        check(close(row["Ad Spend"], R["spend_by_month"][m]), f"Monthly {m}: ad spend matches")
    check(all(row["Net Revenue"] >= 0 for row in mrows), "No month has negative net revenue")
    check(close(sum(r["Net Revenue"] for r in mrows), k["Net revenue"]), "Sum of monthly net revenue = headline net revenue")
    check(sum(r["Orders"] for r in mrows) == k["Orders"], "Sum of monthly orders = headline orders")
    check(all(close(r["AOV"], r["Net Revenue"] / r["Orders"]) for r in mrows), "Monthly AOV = net revenue ÷ orders")

    # ---------------- 3. cohorts
    ws = wb["Cohorts"]
    crows = table(ws, 4)
    sizes = sum(r["Customers"] for r in crows)
    check(sizes == k["Identified customers"],
          f"Cohort sizes sum to identified customers ({sizes:,} vs {k['Identified customers']:,})")
    check(all(close(r["M0"], 1.0) for r in crows), "Every cohort has M0 retention = 100%")
    n_off = sum(1 for key in crows[0] if re.fullmatch(r"M\d+", key))
    check(n_off == len(R["months"]), f"Cohort grid has one column per month offset ({n_off} vs {len(R['months'])})")
    bad = 0
    for r in crows:
        for key, v in r.items():
            if key.startswith("M") and v is not None:
                exp = len(R["cohort_active"][(r["Cohort"], int(key[1:]))]) / R["cohort_size"][r["Cohort"]]
                bad += not close(v, exp, abs_=1e-9)
                bad += not (0 <= v <= 1)
    check(bad == 0, f"All cohort retention cells match raw recompute and lie in [0,1] ({bad} mismatches)")
    cnt_hdr = find_row(ws, "Active customers (counts)") + 1
    cnts = table(ws, cnt_hdr)
    check(sum(r["Customers"] for r in cnts) == sizes, "Cohort count table sizes = retention table sizes")
    check(all((r.get(c) or 0) <= r["Customers"] for r in cnts for c in r if c.startswith("M")),
          "No cohort cell exceeds its cohort size")

    # ---------------- 4. channels
    ws = wb["Channels"]
    chrows = table(ws, 4)
    for r in chrows:
        c = r["Channel"]
        check(close(r["Net Revenue"], R["rev_by_ch"][c]), f"Channel {c}: net revenue matches raw recompute")
        if r["Ad Spend"]:
            check(close(r["ROAS"], r["Net Revenue"] / r["Ad Spend"]), f"Channel {c}: ROAS = revenue ÷ spend")
            check(close(r["CAC"], r["Ad Spend"] / r["New Customers"]), f"Channel {c}: CAC = spend ÷ new customers")
            check(close(r["Ad Spend"], R["spend_by_ch"][c]), f"Channel {c}: spend matches raw recompute")
        else:
            check(r["ROAS"] is None, f"Channel {c}: no spend -> ROAS blank (not divided by zero)")
    check(close(sum(r["Net Revenue"] for r in chrows), k["Net revenue"]), "Channel revenue sums to headline net revenue")
    check(close(sum(r["Ad Spend"] for r in chrows), k["Ad spend (Meta + Google)"]), "Channel spend sums to total spend")
    check(sum(r["New Customers"] for r in chrows) == k["New customers"], "Channel new customers sum to total new customers")
    mhdr = find_row(ws, "Paid channels by month") + 1
    cm = table(ws, mhdr)
    roas_bad = sum(not close(r["ROAS"], r["Net Revenue"] / r["Ad Spend"]) for r in cm)
    raw_bad = sum(not close(r["Ad Spend"], R["spend_by_ch_month"][(r["Channel"], r["Month"])]) or
                  not close(r["Net Revenue"], R["rev_by_ch_month"][(r["Channel"], r["Month"])]) for r in cm)
    check(roas_bad == 0 and len(cm) == 36, f"Monthly paid-channel ROAS = revenue ÷ spend for all {len(cm)} rows")
    check(raw_bad == 0, f"Monthly paid-channel spend & revenue match raw recompute ({raw_bad} mismatches)")

    # ---------------- 5. products
    ws = wb["Products"]
    prows = table(ws, 4)
    check(close(sum(r["Net Revenue"] for r in prows), k["Net revenue"]), "Product net revenue sums to headline")
    check(all(r["Net Revenue"] >= 0 for r in prows), "No product has negative net revenue")
    check(all(close(r["Net Revenue"], R["product_net"][r["SKU"]]) for r in prows), "Every SKU's net revenue matches raw")

    # ---------------- 6. data quality log consistency
    ws = wb["Data Quality Log"]
    dq = table(ws, 4)
    out_row = [r for r in dq if r["Rule"] == "Clean order lines written"][0]
    check(out_row["Rows Affected"] == R["n_clean_lines"],
          f"DQ log clean-row count ({out_row['Rows Affected']:,}) = independent recompute ({R['n_clean_lines']:,})")
    check(close(out_row["Value (USD)"], k["Net revenue"]), "DQ log output value = headline net revenue")

    # ---------------- 7. executive summary ties out
    md = SUMMARY_MD.read_text()
    check("Sample project — synthetic data, not client work" in md, "Executive summary carries the sample-project label")
    mapping = {
        "Net revenue": (f"${k['Net revenue']:,.0f}",),
        "Orders": (f"{k['Orders']:,}",),
        "AOV": (f"${k['AOV']:,.2f}",),
        "Repeat-customer rate (2+ orders)": (f"{k['Repeat-customer rate']:.1%}",),
        "Ad spend (Meta + Google)": (f"${k['Ad spend (Meta + Google)']:,.0f}",),
        "MER (net revenue ÷ ad spend)": (f"{k['MER (net revenue ÷ ad spend)']:.2f}x",),
        "Blended CAC": (f"${k['Blended CAC']:,.2f}",),
        "Meta ROAS / CAC": (f"{k['Meta ROAS']:.2f}x / ${k['Meta CAC']:,.2f}",),
        "Google ROAS / CAC": (f"{k['Google ROAS']:.2f}x / ${k['Google CAC']:,.2f}",),
        "Refund rate": (f"{k['Refund rate']:.1%}",),
        "Gross margin": (f"{k['Gross margin %']:.1%}",),
    }
    md_table = dict(re.findall(r"^\| ([^|]+?) \| ([^|]+?) \|$", md, flags=re.M))
    for label, (expected,) in mapping.items():
        check(md_table.get(label) == expected, f"Exec summary '{label}' = {md_table.get(label)!r}, recomputed {expected!r}")
    check(f"{R['rep_bed']:.1%}" in md and f"{R['rep_oth']:.1%}" in md,
          f"Exec summary bedding-first repeat rates ({R['rep_bed']:.1%} vs {R['rep_oth']:.1%}) match recompute")
    # every supporting figure on the Summary sheet appears in the summary text
    ws = wb["Summary KPIs"]
    sup_row = find_row(ws, "Figure")
    missing = []
    r = sup_row + 1
    while ws.cell(row=r, column=1).value:
        v, f = ws.cell(row=r, column=2).value, ws.cell(row=r, column=2).number_format
        s = (f"{v:.1%}" if "%" in f else f"{v:.2f}x" if "x" in f else f"${v:,.2f}" if ".00" in f else f"${v:,.0f}")
        if s not in md:
            missing.append(f"{ws.cell(row=r, column=1).value}={s}")
        r += 1
    check(not missing, f"All {r - sup_row - 1} supporting figures on Summary sheet appear verbatim in exec summary"
          + (f" (missing: {missing})" if missing else ""))

    # ---------------- report
    fails = [m for ok, m in results if not ok]
    for ok, m in results:
        if not ok or "--verbose" in sys.argv:
            print(("PASS  " if ok else "FAIL  ") + m)
    print(f"\n{len(results) - len(fails)}/{len(results)} checks passed.")
    for name in ["Net revenue", "Orders", "AOV", "Repeat-customer rate", "Meta ROAS", "Google ROAS", "Blended CAC"]:
        print(f"  recomputed {name:22s} {k[name]:,.4f}")
    print("ALL CHECKS PASSED" if not fails else f"{len(fails)} CHECK(S) FAILED")
    return 0 if not fails else 1


if __name__ == "__main__":
    sys.exit(main())
