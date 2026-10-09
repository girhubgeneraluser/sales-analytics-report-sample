"""Build the Northwind performance tables, workbook, charts and executive summary.

Sample project - synthetic data, not client work.

Inputs : data/clean/order_lines_clean.csv, data/clean/ad_spend_clean.csv, data/clean/data_quality_log.csv
Outputs: deliverables/Northwind_Performance_Report.xlsx
         deliverables/executive_summary.md
         deliverables/chart_*.png
         deliverables/kpis.json  (machine-readable headline numbers, used by verify.py)
"""
from __future__ import annotations

import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.drawing.image import Image as XLImage  # noqa: E402
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
CLEAN = ROOT / "data" / "clean"
OUT = ROOT / "deliverables"
SAMPLE_LABEL = "Sample project — synthetic data, not client work"
PAID = ["meta", "google"]

# --------------------------------------------------------------------------- palette (dataviz reference)
SURFACE = "#fcfcfb"
INK = "#0b0b0b"
INK2 = "#52514e"
GRID = "#e4e3df"
BLUE = "#2a78d6"
ORANGE = "#eb6834"
AQUA = "#1baf7a"
SEQ = ["#cde2fb", "#9ec5f4", "#6da7ec", "#3987e5", "#256abf", "#184f95", "#0d366b"]


# =========================================================================== metrics
def load() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    lines = pd.read_csv(CLEAN / "order_lines_clean.csv", keep_default_na=False,
                        dtype={"customer_email": str, "discount_code": str})
    spend = pd.read_csv(CLEAN / "ad_spend_clean.csv")
    dq = pd.read_csv(CLEAN / "data_quality_log.csv")
    return lines, spend, dq


def build_orders(lines: pd.DataFrame) -> pd.DataFrame:
    orders = lines.groupby("order_id").agg(
        created_at_utc=("created_at_utc", "first"), month=("month", "first"),
        customer_email=("customer_email", "first"), country=("billing_country", "first"),
        channel=("source_channel", "first"), gross=("gross_usd", "sum"), discount=("discount_usd", "sum"),
        refund=("refund_usd", "sum"), net=("net_usd", "sum"), cogs=("cogs_usd", "sum"),
        units=("quantity", "sum")).reset_index()
    orders = orders.sort_values(["created_at_utc", "order_id"])
    ident = orders["customer_email"] != ""
    first_month = orders[ident].groupby("customer_email")["month"].min()
    first_order = orders[ident].groupby("customer_email")["order_id"].first()  # sorted by time
    orders["first_month"] = orders["customer_email"].map(first_month)
    orders["is_first_order"] = orders["order_id"].isin(first_order.values)
    return orders


def monthly_table(orders: pd.DataFrame, spend: pd.DataFrame) -> pd.DataFrame:
    g = orders.groupby("month")
    m = pd.DataFrame({
        "orders": g["order_id"].nunique(),
        "gross_sales": g["gross"].sum(),
        "discounts": g["discount"].sum(),
        "refunds": g["refund"].sum(),
        "net_revenue": g["net"].sum(),
        "cogs": g["cogs"].sum(),
    })
    m["aov"] = m["net_revenue"] / m["orders"]
    m["gross_margin_pct"] = (m["net_revenue"] - m["cogs"]) / m["net_revenue"]
    ident = orders[orders["customer_email"] != ""]
    gi = ident.groupby("month")
    m["customers"] = gi["customer_email"].nunique()
    m["new_customers"] = ident[ident["is_first_order"]].groupby("month")["customer_email"].nunique()
    ret = ident[ident["first_month"] < ident["month"]]
    m["returning_customers"] = ret.groupby("month")["customer_email"].nunique()
    m = m.fillna({"new_customers": 0, "returning_customers": 0})
    m["repeat_customer_rate"] = m["returning_customers"] / m["customers"]
    m["ad_spend"] = spend.groupby("month")["spend_usd"].sum()
    m["mer"] = m["net_revenue"] / m["ad_spend"]
    m["blended_cac"] = m["ad_spend"] / m["new_customers"]
    m.index.name = "month"
    return m.reset_index()


def cohort_tables(orders: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    ident = orders[orders["customer_email"] != ""].copy()
    to_idx = lambda s: pd.PeriodIndex(s, freq="M").astype("int64")  # noqa: E731
    ident["k"] = to_idx(ident["month"]) - to_idx(ident["first_month"])
    counts = ident.groupby(["first_month", "k"])["customer_email"].nunique().unstack(fill_value=0)
    counts = counts.reindex(columns=range(orders["month"].nunique()), fill_value=0)  # keep all-zero offsets
    counts.columns = [f"M{c}" for c in counts.columns]
    size = counts["M0"]
    pct = counts.div(size, axis=0)
    # mask cells that are in the future (cohort hasn't been observed that long)
    months = sorted(orders["month"].unique())
    last = len(months) - 1
    for i, cm in enumerate(counts.index):
        age = last - months.index(cm)
        for c in counts.columns:
            if int(c[1:]) > age:
                pct.loc[cm, c] = np.nan
                counts.loc[cm, c] = np.nan
    counts.insert(0, "cohort_size", size)
    pct.insert(0, "cohort_size", size)
    counts.index.name = pct.index.name = "cohort"
    return counts.reset_index(), pct.reset_index()


def channel_tables(orders: pd.DataFrame, spend: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    g = orders.groupby("channel")
    ch = pd.DataFrame({
        "orders": g["order_id"].nunique(),
        "net_revenue": g["net"].sum(),
        "new_customers": orders[orders["is_first_order"]].groupby("channel")["customer_email"].nunique(),
    }).fillna(0)
    ch["ad_spend"] = spend.groupby("channel")["spend_usd"].sum()
    ch["ad_spend"] = ch["ad_spend"].fillna(0)
    ch["roas"] = np.where(ch["ad_spend"] > 0, ch["net_revenue"] / ch["ad_spend"].replace(0, np.nan), np.nan)
    ch["cac"] = np.where(ch["ad_spend"] > 0, ch["ad_spend"] / ch["new_customers"].replace(0, np.nan), np.nan)
    ch["share_of_revenue"] = ch["net_revenue"] / ch["net_revenue"].sum()
    order = ["meta", "google", "email", "organic", "direct", "referral"]
    ch = ch.reindex([c for c in order if c in ch.index] + [c for c in ch.index if c not in order])
    ch.index.name = "channel"

    rows = []
    for c in PAID:
        o = orders[orders["channel"] == c]
        r = o.groupby("month")["net"].sum()
        n = o[o["is_first_order"]].groupby("month")["customer_email"].nunique()
        s = spend[spend["channel"] == c].groupby("month")["spend_usd"].sum()
        t = pd.DataFrame({"ad_spend": s, "net_revenue": r, "new_customers": n}).fillna(0)
        t["roas"] = t["net_revenue"] / t["ad_spend"]
        t["cac"] = t["ad_spend"] / t["new_customers"]
        t["channel"] = c
        rows.append(t)
    chm = pd.concat(rows).reset_index().rename(columns={"index": "month"})
    chm = chm[["month", "channel", "ad_spend", "net_revenue", "roas", "new_customers", "cac"]]
    return ch.reset_index(), chm


def product_table(lines: pd.DataFrame) -> pd.DataFrame:
    g = lines.groupby(["sku", "product_title", "category"])
    p = pd.DataFrame({
        "orders": g["order_id"].nunique(), "units": g["quantity"].sum(), "gross_sales": g["gross_usd"].sum(),
        "discounts": g["discount_usd"].sum(), "refunds": g["refund_usd"].sum(), "net_revenue": g["net_usd"].sum(),
        "cogs": g["cogs_usd"].sum()}).reset_index()
    p["refund_rate"] = p["refunds"] / (p["gross_sales"] - p["discounts"])
    p["gross_profit"] = p["net_revenue"] - p["cogs"]
    p["gross_margin_pct"] = p["gross_profit"] / p["net_revenue"]
    p = p.sort_values("net_revenue", ascending=False).reset_index(drop=True)
    p.insert(0, "rank", range(1, len(p) + 1))
    return p


def headline(orders, monthly, ch, lines) -> dict:
    ident = orders[orders["customer_email"] != ""]
    per_cust = ident.groupby("customer_email")["order_id"].nunique()
    total_spend = monthly["ad_spend"].sum()
    new_total = int(orders["is_first_order"].sum())
    paid = ch.set_index("channel")
    k = {
        "period": f"{monthly['month'].iloc[0]} to {monthly['month'].iloc[-1]}",
        "net_revenue": float(orders["net"].sum()),
        "gross_sales": float(orders["gross"].sum()),
        "discounts": float(orders["discount"].sum()),
        "refunds": float(orders["refund"].sum()),
        "orders": int(orders["order_id"].nunique()),
        "aov": float(orders["net"].sum() / orders["order_id"].nunique()),
        "customers": int(per_cust.size),
        "repeat_customer_rate": float((per_cust >= 2).mean()),
        "ad_spend": float(total_spend),
        "mer": float(orders["net"].sum() / total_spend),
        "blended_cac": float(total_spend / new_total),
        "new_customers": new_total,
        "meta_roas": float(paid.loc["meta", "roas"]),
        "google_roas": float(paid.loc["google", "roas"]),
        "meta_cac": float(paid.loc["meta", "cac"]),
        "google_cac": float(paid.loc["google", "cac"]),
        "refund_rate": float(orders["refund"].sum() / (orders["gross"].sum() - orders["discount"].sum())),
        "gross_margin_pct": float((orders["net"].sum() - orders["cogs"].sum()) / orders["net"].sum()),
    }
    return k


# =========================================================================== charts
def style_axes(ax, title: str, subtitle: str | None = None):
    ax.set_facecolor(SURFACE)
    for s in ["top", "right", "left"]:
        ax.spines[s].set_visible(False)
    ax.spines["bottom"].set_color(INK2)
    ax.tick_params(colors=INK2, labelsize=10, length=0)
    ax.grid(axis="y", color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)
    ax.set_title(title, loc="left", fontsize=14, color=INK, fontweight="bold", pad=24 if subtitle else 10)
    if subtitle:
        ax.text(0, 1.02, subtitle, transform=ax.transAxes, fontsize=10, color=INK2, va="bottom")


def footer(fig):
    fig.text(0.01, 0.01, SAMPLE_LABEL, fontsize=8, color=INK2)


def month_labels(months):
    return [pd.Period(m).strftime("%b\n%y") if i % 3 == 0 else pd.Period(m).strftime("%b")
            for i, m in enumerate(months)]


def chart_revenue(monthly: pd.DataFrame, path: Path):
    fig, ax = plt.subplots(figsize=(11, 5.2), facecolor=SURFACE)
    x = np.arange(len(monthly))
    vals = monthly["net_revenue"] / 1000
    ax.bar(x, vals, color=BLUE, width=0.72, edgecolor=SURFACE, linewidth=2)
    ax.set_xticks(x, month_labels(monthly["month"]))
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: f"${v:,.0f}k"))
    peak = int(vals.idxmax())
    for i in sorted({0, peak, len(vals) - 1}):
        ax.annotate(f"${vals[i]:,.0f}k", (x[i], vals[i]), xytext=(0, 4), textcoords="offset points",
                    ha="center", fontsize=10, color=INK)
    style_axes(ax, "Net revenue by month (USD)",
               "After discounts and refunds. November peaks are Black Friday / Cyber Monday.")
    footer(fig)
    fig.tight_layout(rect=(0, 0.03, 1, 1))
    fig.savefig(path, dpi=150, facecolor=SURFACE)
    plt.close(fig)


def chart_roas(chm: pd.DataFrame, breakeven: float, path: Path):
    fig, ax = plt.subplots(figsize=(11, 5.2), facecolor=SURFACE)
    months = sorted(chm["month"].unique())
    x = np.arange(len(months))
    for c, col, label in (("meta", ORANGE, "Meta"), ("google", BLUE, "Google")):
        s = chm[chm["channel"] == c].set_index("month").reindex(months)["roas"]
        ax.plot(x, s.values, color=col, linewidth=2.2, marker="o", markersize=6,
                markeredgecolor=SURFACE, markeredgewidth=1.5, label=label)
        ax.annotate(f"{label} {s.values[-1]:.1f}x", (x[-1], s.values[-1]), xytext=(8, 0),
                    textcoords="offset points", va="center", fontsize=10, color=INK)
    ax.axhline(breakeven, color=INK2, linestyle="--", linewidth=1.2)
    ax.text(0, breakeven, f"  break-even ROAS {breakeven:.2f}x (1 ÷ gross margin)", va="bottom",
            fontsize=9, color=INK2)
    ax.set_xticks(x, month_labels(months))
    ax.set_xlim(-0.5, len(months) + 1.3)
    ax.set_ylim(0, None)
    ax.yaxis.set_major_locator(matplotlib.ticker.MultipleLocator(1))
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: f"{v:.0f}x"))
    ax.legend(loc="upper right", frameon=False, fontsize=10, ncol=2, bbox_to_anchor=(1, 1.12))
    style_axes(ax, "Paid channel ROAS by month",
               "Last-click net revenue ÷ spend. Meta efficiency has fallen as spend scaled.")
    footer(fig)
    fig.tight_layout(rect=(0, 0.03, 1, 1))
    fig.savefig(path, dpi=150, facecolor=SURFACE)
    plt.close(fig)


def chart_cohorts(pct: pd.DataFrame, path: Path, max_k: int = 12):
    cols = [f"M{k}" for k in range(1, max_k + 1) if f"M{k}" in pct.columns]
    data = pct.set_index("cohort")[cols].astype(float) * 100
    data = data.dropna(how="all")  # newest cohort has no follow-up month yet
    sizes = pct.set_index("cohort").loc[data.index, "cohort_size"]
    fig, ax = plt.subplots(figsize=(11, 6.4), facecolor=SURFACE)
    cmap = matplotlib.colors.LinearSegmentedColormap.from_list("seq", SEQ)
    vmax = np.nanpercentile(data.values, 98)
    im = ax.imshow(np.ma.masked_invalid(data.values), cmap=cmap, aspect="auto", vmin=0, vmax=vmax)
    for i in range(data.shape[0]):
        for j in range(data.shape[1]):
            v = data.values[i, j]
            if not np.isnan(v):
                ax.text(j, i, f"{v:.1f}", ha="center", va="center", fontsize=8,
                        color="white" if v > vmax * 0.55 else INK)
    ax.set_xticks(range(len(cols)), [c.replace("M", "+") for c in cols])
    ax.set_yticks(range(data.shape[0]),
                  [f"{pd.Period(c).strftime('%b %Y')}  (n={int(n):,})" for c, n in
                   zip(data.index, sizes)])
    ax.set_xlabel("Months after first order", color=INK2)
    ax.set_facecolor(SURFACE)
    for s in ax.spines.values():
        s.set_visible(False)
    ax.tick_params(colors=INK2, labelsize=9, length=0)
    cb = fig.colorbar(im, ax=ax, fraction=0.025, pad=0.02)
    cb.set_label("% of cohort ordering again", color=INK2)
    cb.outline.set_visible(False)
    cb.ax.tick_params(colors=INK2, labelsize=9)
    ax.set_title("Cohort retention: % of first-time customers who order again", loc="left",
                 fontsize=14, fontweight="bold", color=INK, pad=12)
    footer(fig)
    fig.tight_layout(rect=(0, 0.03, 1, 1))
    fig.savefig(path, dpi=150, facecolor=SURFACE)
    plt.close(fig)


def chart_products(p: pd.DataFrame, path: Path):
    fig, axes = plt.subplots(1, 2, figsize=(13, 5.6), facecolor=SURFACE)
    top = p.nlargest(10, "net_revenue").iloc[::-1]
    axes[0].barh(top["product_title"], top["net_revenue"] / 1000, color=BLUE, height=0.7,
                 edgecolor=SURFACE, linewidth=2)
    for y, v in enumerate(top["net_revenue"] / 1000):
        axes[0].text(v, y, f" ${v:,.0f}k", va="center", fontsize=9, color=INK)
    axes[0].set_xlim(0, top["net_revenue"].max() / 1000 * 1.2)
    axes[0].xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: f"${v:,.0f}k"))
    axes[0].xaxis.set_major_locator(matplotlib.ticker.MultipleLocator(50))
    style_axes(axes[0], "Top 10 products by net revenue")
    axes[0].grid(axis="y", visible=False)
    axes[0].grid(axis="x", color=GRID)

    rr = p[p["orders"] >= 50].nlargest(10, "refund_rate").iloc[::-1]
    avg = p["refunds"].sum() / (p["gross_sales"].sum() - p["discounts"].sum())
    axes[1].barh(rr["product_title"], rr["refund_rate"] * 100, color=ORANGE, height=0.7,
                 edgecolor=SURFACE, linewidth=2)
    for y, v in enumerate(rr["refund_rate"] * 100):
        axes[1].text(v, y, f" {v:.1f}%", va="center", fontsize=9, color=INK)
    axes[1].axvline(avg * 100, color=INK2, linestyle="--", linewidth=1.2)
    axes[1].text(avg * 100, -0.9, f" store avg {avg*100:.1f}%", fontsize=9, color=INK2, va="center")
    axes[1].set_xlim(0, rr["refund_rate"].max() * 100 * 1.25)
    axes[1].set_ylim(-1.3, len(rr) - 0.5)
    axes[1].xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: f"{v:.0f}%"))
    style_axes(axes[1], "Highest refund rates (50+ orders)")
    axes[1].grid(axis="y", visible=False)
    axes[1].grid(axis="x", color=GRID)
    footer(fig)
    fig.tight_layout(rect=(0, 0.03, 1, 1))
    fig.savefig(path, dpi=150, facecolor=SURFACE)
    plt.close(fig)


# =========================================================================== workbook
HDR_FILL = PatternFill("solid", fgColor="1C5CAB")
HDR_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="0B0B0B")
NOTE_FONT = Font(italic=True, size=9, color="52514E")
THIN = Border(bottom=Side(style="thin", color="E4E3DF"))
FMT = {"usd": '"$"#,##0', "usd2": '"$"#,##0.00', "int": "#,##0", "pct": "0.0%", "x": '0.00"x"', "txt": "@"}


def write_table(ws, df: pd.DataFrame, start_row: int, formats: dict[str, str], headers: dict[str, str] | None = None,
                start_col: int = 1) -> tuple[int, int]:
    headers = headers or {}
    for j, c in enumerate(df.columns):
        cell = ws.cell(row=start_row, column=start_col + j, value=headers.get(c, c.replace("_", " ").title()))
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(row):
            if isinstance(v, float) and np.isnan(v):
                v = None
            elif isinstance(v, (np.integer,)):
                v = int(v)
            elif isinstance(v, (np.floating,)):
                v = float(v)
            cell = ws.cell(row=start_row + i, column=start_col + j, value=v)
            f = formats.get(df.columns[j])
            if f:
                cell.number_format = FMT[f]
            cell.border = THIN
    ws.row_dimensions[start_row].height = 32
    return start_row + len(df), start_col + len(df.columns) - 1


def sheet_header(ws, title: str, note: str):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"{SAMPLE_LABEL}. {note}"
    ws["A2"].font = NOTE_FONT
    ws.sheet_view.showGridLines = False


def widths(ws, w: dict[str, float] | float, n: int = 0):
    if isinstance(w, dict):
        for k, v in w.items():
            ws.column_dimensions[k].width = v
    else:
        for i in range(1, n + 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def build_workbook(k, monthly, counts, pct, ch, chm, prod, dq, charts, path: Path):
    wb = Workbook()

    # ---------------- Summary KPIs
    ws = wb.active
    ws.title = "Summary KPIs"
    sheet_header(ws, "Northwind Home Goods — Performance Report", f"Period {k['period']}, all values USD.")
    rows = [
        ("Net revenue", k["net_revenue"], "usd", "Gross sales − discounts − refunds"),
        ("Gross sales", k["gross_sales"], "usd", "Quantity × unit price"),
        ("Discounts", k["discounts"], "usd", ""),
        ("Refunds", k["refunds"], "usd", ""),
        ("Orders", k["orders"], "int", "Distinct valid orders"),
        ("AOV", k["aov"], "usd2", "Net revenue ÷ orders"),
        ("Identified customers", k["customers"], "int", "Unique emails"),
        ("Repeat-customer rate", k["repeat_customer_rate"], "pct", "Share of customers with 2+ orders"),
        ("New customers", k["new_customers"], "int", "First-ever order in period"),
        ("Ad spend (Meta + Google)", k["ad_spend"], "usd", ""),
        ("MER (net revenue ÷ ad spend)", k["mer"], "x", "Blended marketing efficiency"),
        ("Blended CAC", k["blended_cac"], "usd2", "Ad spend ÷ all new customers"),
        ("Meta ROAS", k["meta_roas"], "x", "Last-click"),
        ("Google ROAS", k["google_roas"], "x", "Last-click"),
        ("Meta CAC", k["meta_cac"], "usd2", "Meta spend ÷ Meta-acquired new customers"),
        ("Google CAC", k["google_cac"], "usd2", "Google spend ÷ Google-acquired new customers"),
        ("Refund rate", k["refund_rate"], "pct", "Refunds ÷ (gross − discounts)"),
        ("Gross margin %", k["gross_margin_pct"], "pct", "(Net revenue − COGS) ÷ net revenue"),
    ]
    for j, h in enumerate(["KPI", "Value", "Definition"], start=1):
        c = ws.cell(row=4, column=j, value=h)
        c.fill, c.font = HDR_FILL, HDR_FONT
    for i, (name, val, f, d) in enumerate(rows, start=5):
        ws.cell(row=i, column=1, value=name).font = Font(bold=True)
        c = ws.cell(row=i, column=2, value=val)
        c.number_format = FMT[f]
        ws.cell(row=i, column=3, value=d).font = Font(color="52514E")
        for j in range(1, 4):
            ws.cell(row=i, column=j).border = THIN
    # latest month comparison
    r0 = 5 + len(rows) + 2
    ws.cell(row=r0 - 1, column=1, value="Latest month vs prior month and same month last year").font = Font(bold=True, size=12)
    mi = monthly.set_index("month")
    last, prev, yoy = mi.index[-1], mi.index[-2], mi.index[-13]
    comp = pd.DataFrame({
        "metric": ["Net revenue", "Orders", "AOV", "Repeat-customer rate", "Ad spend", "MER", "Blended CAC"],
        last: [mi.loc[last, c] for c in ["net_revenue", "orders", "aov", "repeat_customer_rate", "ad_spend", "mer", "blended_cac"]],
        prev: [mi.loc[prev, c] for c in ["net_revenue", "orders", "aov", "repeat_customer_rate", "ad_spend", "mer", "blended_cac"]],
        yoy: [mi.loc[yoy, c] for c in ["net_revenue", "orders", "aov", "repeat_customer_rate", "ad_spend", "mer", "blended_cac"]],
    })
    comp["vs_prior_month"] = comp[last] / comp[prev] - 1
    comp["vs_last_year"] = comp[last] / comp[yoy] - 1
    end, _ = write_table(ws, comp, r0, {"vs_prior_month": "pct", "vs_last_year": "pct"},
                         {"metric": "Metric", "vs_prior_month": "Δ vs prior month (relative)", "vs_last_year": "Δ vs last year (relative)"})
    fmts = ["usd", "int", "usd2", "pct", "usd", "x", "usd2"]
    for i, f in enumerate(fmts):
        for j in (2, 3, 4):
            ws.cell(row=r0 + 1 + i, column=j).number_format = FMT[f]
    # supporting figures quoted in the executive summary
    r1 = end + 3
    ws.cell(row=r1 - 1, column=1, value="Supporting figures quoted in the executive summary").font = Font(bold=True, size=12)
    sup = [
        ("Meta spend, first 6 months", k["meta_spend_first6"], "usd", "Channels sheet, monthly block"),
        ("Meta spend, last 6 months", k["meta_spend_last6"], "usd", ""),
        ("Meta ROAS, first 6 months", k["meta_roas_first6"], "x", ""),
        ("Meta ROAS, last 6 months", k["meta_roas_last6"], "x", ""),
        ("Meta CAC, first 6 months", k["meta_cac_first6"], "usd2", ""),
        ("Meta CAC, last 6 months", k["meta_cac_last6"], "usd2", ""),
        ("Google ROAS, last 6 months", k["google_roas_last6"], "x", ""),
        ("Google CAC, last 6 months", k["google_cac_last6"], "usd2", ""),
        ("Break-even ROAS (1 ÷ gross margin)", k["breakeven_roas"], "x", ""),
        ("Month-1 retention (size-weighted)", k["m1_retention"], "pct", "Cohorts sheet, M1 column"),
        ("Repeat rate, Bedding in first order", k["repeat_bedding_first"], "pct", "Customers with 2+ orders"),
        ("Repeat rate, no Bedding in first order", k["repeat_other_first"], "pct", ""),
        (f"Refund rate, {k['worst_refund_sku']}", k["worst_refund_rate"], "pct", "Products sheet"),
        (f"Refunds, {k['worst_refund_sku']}", k["worst_refund_usd"], "usd", ""),
        (f"Refunds above store-average rate, {k['worst_refund_sku']}", k["worst_excess_refunds"], "usd", ""),
        (f"Discount rate, {k['bfcm_month']}", k["nov_discount_rate"], "pct", "Discounts ÷ gross sales"),
        ("Discount rate, all other months", k["other_discount_rate"], "pct", ""),
        (f"Gross margin, {k['bfcm_month']}", k["bfcm_gross_margin"], "pct", "Monthly sheet"),
    ]
    for j, h in enumerate(["Figure", "Value", "Source"], start=1):
        c = ws.cell(row=r1, column=j, value=h)
        c.fill, c.font = HDR_FILL, HDR_FONT
    for i, (name, val, f, d) in enumerate(sup, start=r1 + 1):
        ws.cell(row=i, column=1, value=name)
        c = ws.cell(row=i, column=2, value=val)
        c.number_format = FMT[f]
        ws.cell(row=i, column=3, value=d).font = Font(color="52514E")
        for j in range(1, 4):
            ws.cell(row=i, column=j).border = THIN
    widths(ws, {"A": 44, "B": 16, "C": 46, "D": 14, "E": 18, "F": 16})
    ws.freeze_panes = "A5"

    # ---------------- Monthly
    ws = wb.create_sheet("Monthly")
    sheet_header(ws, "Monthly performance", "Months in store time (America/New_York).")
    mf = {"orders": "int", "gross_sales": "usd", "discounts": "usd", "refunds": "usd", "net_revenue": "usd",
          "cogs": "usd", "aov": "usd2", "gross_margin_pct": "pct", "customers": "int", "new_customers": "int",
          "returning_customers": "int", "repeat_customer_rate": "pct", "ad_spend": "usd", "mer": "x",
          "blended_cac": "usd2"}
    end, last_col = write_table(ws, monthly, 4, mf, {"aov": "AOV", "cogs": "COGS", "mer": "MER",
                                                     "blended_cac": "Blended CAC", "gross_margin_pct": "Gross Margin %"})
    tr = end + 1
    ws.cell(row=tr, column=1, value="Total").font = Font(bold=True)
    for j, c in enumerate(monthly.columns[1:], start=2):
        if c in ("orders", "gross_sales", "discounts", "refunds", "net_revenue", "cogs", "ad_spend", "new_customers"):
            col = get_column_letter(j)
            cell = ws.cell(row=tr, column=j, value=f"=SUM({col}5:{col}{end})")
            cell.number_format = FMT[mf[c]]
            cell.font = Font(bold=True)
    ws.cell(row=tr + 1, column=1, value="Total row uses live SUM formulas; AOV / rates are not summable — see Summary KPIs.").font = NOTE_FONT
    ws.conditional_formatting.add(f"F5:F{end}", DataBarRule(start_type="min", end_type="max", color="6DA7EC"))
    widths(ws, 14, last_col)
    ws.freeze_panes = "B5"

    # ---------------- Cohorts
    ws = wb.create_sheet("Cohorts")
    sheet_header(ws, "Customer cohort retention",
                 "Cohort = month of first order. Cell = % of cohort ordering in month +k. Blank = not yet observable.")
    ws["A3"] = "Retention %"
    ws["A3"].font = Font(bold=True, size=12)
    pf = {c: "pct" for c in pct.columns if c.startswith("M")}
    pf["cohort_size"] = "int"
    end, last_col = write_table(ws, pct, 4, pf, {"cohort": "Cohort", "cohort_size": "Customers"})
    first_m = get_column_letter(4)  # M1 column (A cohort, B size, C M0, D M1)
    ws.conditional_formatting.add(
        f"{first_m}5:{get_column_letter(last_col)}{end}",
        ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                       mid_type="percentile", mid_value=50, mid_color="9EC5F4",
                       end_type="max", end_color="1C5CAB"))
    r2 = end + 3
    ws.cell(row=r2 - 1, column=1, value="Active customers (counts)").font = Font(bold=True, size=12)
    cf = {c: "int" for c in counts.columns if c != "cohort"}
    end2, _ = write_table(ws, counts, r2, cf, {"cohort": "Cohort", "cohort_size": "Customers"})
    tr = end2 + 1
    ws.cell(row=tr, column=1, value="Total").font = Font(bold=True)
    c = ws.cell(row=tr, column=2, value=f"=SUM(B{r2 + 1}:B{end2})")
    c.number_format, c.font = FMT["int"], Font(bold=True)
    ws.cell(row=tr + 1, column=1,
            value="Cohort sizes sum to the total number of identified customers (checked in verify.py).").font = NOTE_FONT
    widths(ws, 8.5, last_col)
    ws.column_dimensions["A"].width = 11
    ws.column_dimensions["B"].width = 11
    ws.freeze_panes = "C5"

    # ---------------- Channels
    ws = wb.create_sheet("Channels")
    sheet_header(ws, "Channel performance",
                 "Last-click attribution from Shopify source. ROAS = net revenue ÷ spend; CAC = spend ÷ new customers.")
    cfm = {"orders": "int", "net_revenue": "usd", "new_customers": "int", "ad_spend": "usd", "roas": "x",
           "cac": "usd2", "share_of_revenue": "pct"}
    end, _ = write_table(ws, ch, 4, cfm, {"roas": "ROAS", "cac": "CAC"})
    r2 = end + 3
    ws.cell(row=r2 - 1, column=1, value="Paid channels by month").font = Font(bold=True, size=12)
    end2, last_col = write_table(ws, chm, r2, {"ad_spend": "usd", "net_revenue": "usd", "roas": "x",
                                               "new_customers": "int", "cac": "usd2"}, {"roas": "ROAS", "cac": "CAC"})
    ws.conditional_formatting.add(f"E{r2 + 1}:E{end2}", ColorScaleRule(
        start_type="min", start_color="F8D3C3", mid_type="percentile", mid_value=50, mid_color="FFFFFF",
        end_type="max", end_color="9EC5F4"))
    widths(ws, 15, 8)
    ws.freeze_panes = "A5"

    # ---------------- Products
    ws = wb.create_sheet("Products")
    sheet_header(ws, "Product performance", "Ranked by net revenue. COGS assumes refunded units are not restocked.")
    pfm = {"orders": "int", "units": "int", "gross_sales": "usd", "discounts": "usd", "refunds": "usd",
           "net_revenue": "usd", "cogs": "usd", "refund_rate": "pct", "gross_profit": "usd", "gross_margin_pct": "pct"}
    end, last_col = write_table(ws, prod, 4, pfm, {"sku": "SKU", "cogs": "COGS", "gross_margin_pct": "Gross Margin %"})
    ws.conditional_formatting.add(f"I5:I{end}", DataBarRule(start_type="min", end_type="max", color="6DA7EC"))
    ws.conditional_formatting.add(f"K5:K{end}", ColorScaleRule(start_type="min", start_color="FFFFFF",
                                                              end_type="max", end_color="EB6834"))
    widths(ws, 13, last_col)
    ws.column_dimensions["C"].width = 32
    ws.freeze_panes = "D5"

    # ---------------- Charts
    ws = wb.create_sheet("Charts")
    sheet_header(ws, "Charts", "Same PNGs as in the deliverables folder.")
    row = 4
    for p in charts:
        img = XLImage(str(p))
        img.width, img.height = img.width * 0.55, img.height * 0.55
        ws.add_image(img, f"A{row}")
        row += int(img.height / 20) + 3

    # ---------------- Data Quality Log
    ws = wb.create_sheet("Data Quality Log")
    sheet_header(ws, "Data quality log", "Every cleaning rule applied by src/clean.py, in order. Nothing is dropped silently.")
    end, last_col = write_table(ws, dq, 4, {"rows_affected": "int", "orders_affected": "int", "value_usd": "usd"},
                                {"value_usd": "Value (USD)"})
    widths(ws, {"A": 6, "B": 10, "C": 62, "D": 12, "E": 12, "F": 14, "G": 80})
    for r in range(5, end + 1):
        ws.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=7).alignment = Alignment(wrap_text=True, vertical="top")

    # ---------------- README
    ws = wb.create_sheet("README")
    sheet_header(ws, "README — how this workbook works and how to refresh it", "")
    text = [
        ("What this is", "Monthly performance report for Northwind Home Goods (a fictional brand). Built from a Shopify "
                         "order export and Meta/Google daily ad spend. All data is synthetic."),
        ("How to refresh (ops)", "1) Export orders from Shopify (Orders → Export → all orders, CSV, one row per line item) "
                                 "and save as data/raw/shopify_orders_export.csv."),
        ("", "2) Export daily spend from Meta Ads Manager and Google Ads (date, campaign, spend) and append to "
             "data/raw/ad_spend_daily.csv with a 'channel' column (meta / google)."),
        ("", "3) Add the new month's CAD→USD rate to data/raw/fx_rates_monthly.csv (finance provides this). "
             "The pipeline stops with an error if a month is missing."),
        ("", "4) From the project folder run:  python3 src/run_all.py   (≈10 seconds). It cleans, rebuilds this "
             "workbook, the charts and the executive summary, then runs the QA checks."),
        ("", "5) Check the QA output says ALL CHECKS PASSED, and skim the Data Quality Log sheet for unusual counts."),
        ("Do not edit by hand", "This file is regenerated on every refresh — manual edits will be overwritten. "
                                "Put commentary in a separate file."),
        ("Definitions", "Net revenue = gross sales − discounts − refunds (excl. shipping & tax). AOV = net revenue ÷ orders. "
                        "Repeat-customer rate (monthly) = returning customers ÷ customers that month; (overall) = share of "
                        "customers with 2+ orders. ROAS = last-click net revenue ÷ spend. CAC = spend ÷ new customers "
                        "acquired by that channel. MER = net revenue ÷ total ad spend."),
        ("Time zone", "Shopify timestamps arrive in mixed formats; all are converted to UTC then bucketed into months "
                      "in America/New_York (store time). Naive timestamps are assumed to be store time."),
        ("Currency", "CAD orders are converted to USD with the monthly finance rate in fx_rates_monthly.csv."),
        ("Customers", "A customer is a unique, normalised email. Orders without an email count toward revenue and "
                      "orders but not toward customer metrics."),
        ("Sheets", "Summary KPIs · Monthly · Cohorts · Channels · Products · Charts · Data Quality Log · README"),
    ]
    for i, (a, b) in enumerate(text, start=4):
        ws.cell(row=i, column=1, value=a).font = Font(bold=True)
        c = ws.cell(row=i, column=2, value=b)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 48 if len(b) > 110 else 32
    widths(ws, {"A": 22, "B": 120})

    wb.save(path)


# =========================================================================== executive summary
def fmt_usd(v, d=0):
    return f"${v:,.{d}f}"


def write_summary(k, monthly, pct, ch, chm, prod, orders, lines, dq, path: Path):
    n_neg = int(dq.loc[dq["rule"].str.startswith("Non-positive quantity"), "rows_affected"].iloc[0])
    n_noemail = int(dq.loc[dq["rule"].str.startswith("Missing email, not recoverable"), "orders_affected"].iloc[0])
    mi = monthly.set_index("month")
    last = mi.index[-1]
    yoy = mi.index[-13]
    mlast = pd.Period(last).strftime("%b %Y")
    myoy = pd.Period(yoy).strftime("%b %Y")
    chi = ch.set_index("channel")

    # Meta: first 6 vs last 6 months
    meta = chm[chm["channel"] == "meta"].set_index("month")
    goog = chm[chm["channel"] == "google"].set_index("month")
    h1 = meta.iloc[:6]
    h2 = meta.iloc[-6:]
    meta_roas_h1 = h1["net_revenue"].sum() / h1["ad_spend"].sum()
    meta_roas_h2 = h2["net_revenue"].sum() / h2["ad_spend"].sum()
    meta_cac_h1 = h1["ad_spend"].sum() / h1["new_customers"].sum()
    meta_cac_h2 = h2["ad_spend"].sum() / h2["new_customers"].sum()
    meta_spend_h1 = h1["ad_spend"].sum()
    meta_spend_h2 = h2["ad_spend"].sum()
    g2 = goog.iloc[-6:]
    goog_roas_h2 = g2["net_revenue"].sum() / g2["ad_spend"].sum()
    goog_cac_h2 = g2["ad_spend"].sum() / g2["new_customers"].sum()
    breakeven = 1 / k["gross_margin_pct"]

    # product refunds
    # product with the largest refund dollars ABOVE the store-average refund rate (biggest $ problem)
    pp = prod[prod["orders"] >= 50].copy()
    pp["excess_refunds"] = pp["refunds"] - k["refund_rate"] * (pp["gross_sales"] - pp["discounts"])
    worst = pp.sort_values("excess_refunds", ascending=False).iloc[0]
    worst_ref = worst["refunds"]
    worst_excess = worst["excess_refunds"]

    # retention: M1 avg (cohorts with M1 observed) and bedding-first effect
    m1 = pct["M1"].dropna()
    m1_avg = (pct.loc[m1.index, "M1"] * pct.loc[m1.index, "cohort_size"]).sum() / pct.loc[m1.index, "cohort_size"].sum()
    ident = orders[orders["customer_email"] != ""]
    first = ident[ident["is_first_order"]][["order_id", "customer_email"]]
    bed_orders = set(lines.loc[lines["category"] == "Bedding", "order_id"])
    first["bedding_first"] = first["order_id"].isin(bed_orders)
    n_orders = ident.groupby("customer_email")["order_id"].nunique()
    first["repeat"] = first["customer_email"].map(n_orders) >= 2
    rep_bed = first.loc[first["bedding_first"], "repeat"].mean()
    rep_other = first.loc[~first["bedding_first"], "repeat"].mean()

    # BFCM discount depth
    nov = [m for m in mi.index if m.endswith("-11")]
    nov_m = nov[-1]
    disc_rate_nov = mi.loc[nov_m, "discounts"] / mi.loc[nov_m, "gross_sales"]
    disc_rate_rest = (mi["discounts"].sum() - mi.loc[nov_m, "discounts"]) / (mi["gross_sales"].sum() - mi.loc[nov_m, "gross_sales"])
    gm_nov = mi.loc[nov_m, "gross_margin_pct"]

    top3 = prod.head(3)
    top3_share = top3["net_revenue"].sum() / prod["net_revenue"].sum()

    k.update({
        "meta_roas_first6": meta_roas_h1, "meta_roas_last6": meta_roas_h2, "meta_cac_first6": meta_cac_h1,
        "meta_cac_last6": meta_cac_h2, "google_roas_last6": goog_roas_h2, "google_cac_last6": goog_cac_h2,
        "breakeven_roas": breakeven, "worst_refund_sku": worst["sku"], "worst_refund_rate": float(worst["refund_rate"]),
        "worst_refund_usd": float(worst_ref), "worst_excess_refunds": float(worst_excess), "m1_retention": float(m1_avg), "repeat_bedding_first": float(rep_bed),
        "repeat_other_first": float(rep_other), "nov_discount_rate": float(disc_rate_nov),
        "other_discount_rate": float(disc_rate_rest), "latest_month": last,
        "meta_spend_first6": float(meta_spend_h1), "meta_spend_last6": float(meta_spend_h2),
        "bfcm_month": nov_m, "bfcm_gross_margin": float(gm_nov),
        "latest_net_revenue": float(mi.loc[last, "net_revenue"]), "yoy_net_revenue": float(mi.loc[yoy, "net_revenue"]),
    })

    md = f"""# Northwind Home Goods — Performance Summary ({pd.Period(monthly['month'].iloc[0]).strftime('%b %Y')} – {mlast})

> **{SAMPLE_LABEL}.** Northwind Home Goods is fictional. Figures tie out to `Northwind_Performance_Report.xlsx` (sheet *Summary KPIs* unless noted).

**Bottom line:** the business grew — {mlast} net revenue was {fmt_usd(mi.loc[last, 'net_revenue'])}, {mi.loc[last, 'net_revenue'] / mi.loc[yoy, 'net_revenue'] - 1:+.0%} vs {myoy} — but growth is increasingly bought through Meta at worsening efficiency, and too few first-time buyers come back.

## Headline numbers (18 months)

| KPI | Value |
|---|---|
| Net revenue | {fmt_usd(k['net_revenue'])} |
| Orders | {k['orders']:,} |
| AOV | {fmt_usd(k['aov'], 2)} |
| Repeat-customer rate (2+ orders) | {k['repeat_customer_rate']:.1%} |
| Ad spend (Meta + Google) | {fmt_usd(k['ad_spend'])} |
| MER (net revenue ÷ ad spend) | {k['mer']:.2f}x |
| Blended CAC | {fmt_usd(k['blended_cac'], 2)} |
| Meta ROAS / CAC | {k['meta_roas']:.2f}x / {fmt_usd(k['meta_cac'], 2)} |
| Google ROAS / CAC | {k['google_roas']:.2f}x / {fmt_usd(k['google_cac'], 2)} |
| Refund rate | {k['refund_rate']:.1%} |
| Gross margin | {k['gross_margin_pct']:.1%} |

## What changed

- **Meta is getting more expensive.** Spend rose from {fmt_usd(meta_spend_h1)} (first 6 months) to {fmt_usd(meta_spend_h2)} (last 6), while last-click ROAS fell from {meta_roas_h1:.2f}x to {meta_roas_h2:.2f}x and CAC climbed from {fmt_usd(meta_cac_h1, 2)} to {fmt_usd(meta_cac_h2, 2)}. Google held at {goog_roas_h2:.2f}x ROAS / {fmt_usd(goog_cac_h2, 2)} CAC over the same last 6 months. Break-even ROAS at our {k['gross_margin_pct']:.0%} gross margin is {breakeven:.2f}x. *(Channels sheet)*
- **Retention is the weak link.** Only {m1_avg:.1%} of new customers order again the month after their first purchase, and {k['repeat_customer_rate']:.1%} ever place a second order. Customers whose first order included Bedding repeat at {rep_bed:.1%} vs {rep_other:.1%} for everyone else. *(Cohorts sheet)*
- **One product drives a disproportionate share of refunds.** {worst['product_title']} ({worst['sku']}) has a {worst['refund_rate']:.1%} refund rate vs a {k['refund_rate']:.1%} store average — {fmt_usd(worst_ref)} refunded over the period, {fmt_usd(worst_excess)} more than if it refunded at the store average. It is the #{int(worst['rank'])} product by net revenue. *(Products sheet)*
- **Black Friday discounting is deep.** In {pd.Period(nov_m).strftime('%b %Y')} discounts were {disc_rate_nov:.1%} of gross sales (vs {disc_rate_rest:.1%} the rest of the period) and gross margin dipped to {gm_nov:.1%}. *(Monthly sheet)*

## Recommendations

1. **Cap Meta prospecting and move ~20% of its budget to Google Shopping + retention.** Meta's last-6-month ROAS of {meta_roas_h2:.2f}x is now close to the {breakeven:.2f}x break-even, and its CAC of {fmt_usd(meta_cac_h2, 2)} is {meta_cac_h2 / goog_cac_h2 - 1:.0%} above Google's. Set a Meta CAC ceiling of {fmt_usd(goog_cac_h2 * 1.1, 0)} and review weekly; run a 4-week geo holdout to measure Meta's true incrementality before scaling again.
2. **Build a 30/60-day second-purchase program.** With month-1 retention at {m1_avg:.1%}, a post-purchase email/SMS flow (care guide → complementary product → time-limited VIP offer) is the cheapest growth lever. Target: lift the repeat-customer rate from {k['repeat_customer_rate']:.1%} to {k['repeat_customer_rate'] + 0.03:.1%} within two quarters.
3. **Lead acquisition with Bedding.** Bedding-first customers repeat {rep_bed / rep_other:.1f}x as often. Shift prospecting creative and the welcome offer toward sheet sets and pillows.
4. **Fix or pull {worst['product_title']}.** Pull return reasons from Shopify for {worst['sku']}; if it's sizing/colour, update the PDP (dimensions, real-light photos). If refund rate doesn't drop below 10% within 60 days, pause paid promotion of the SKU.
5. **Make BFCM shallower.** Test a 15–20% tiered offer (spend-more-save-more) instead of a flat 25% code; every 5 points of discount on November's gross sales is ≈{fmt_usd(mi.loc[nov_m, 'gross_sales'] * 0.05)} of margin.

*Method & caveats: last-click attribution from Shopify order source; CAD converted at monthly finance rates; {n_neg} lines with negative quantities quarantined and {n_noemail} orders without an email excluded from customer metrics (see Data Quality Log). Full definitions in the workbook README sheet.*
"""
    path.write_text(md)
    return k


# =========================================================================== main
def main() -> None:
    OUT.mkdir(exist_ok=True)
    lines, spend, dq = load()
    orders = build_orders(lines)
    monthly = monthly_table(orders, spend)
    counts, pct = cohort_tables(orders)
    ch, chm = channel_tables(orders, spend)
    prod = product_table(lines)
    k = headline(orders, monthly, ch, lines)

    charts = [OUT / "chart_1_monthly_revenue.png", OUT / "chart_2_channel_roas.png",
              OUT / "chart_3_cohort_retention.png", OUT / "chart_4_products.png"]
    chart_revenue(monthly, charts[0])
    chart_roas(chm, 1 / k["gross_margin_pct"], charts[1])
    chart_cohorts(pct, charts[2])
    chart_products(prod, charts[3])

    k = write_summary(k, monthly, pct, ch, chm, prod, orders, lines, dq, OUT / "executive_summary.md")
    build_workbook(k, monthly, counts, pct, ch, chm, prod, dq, charts, OUT / "Northwind_Performance_Report.xlsx")
    (OUT / "kpis.json").write_text(json.dumps(k, indent=2, default=float))
    for key in ["net_revenue", "orders", "aov", "customers", "repeat_customer_rate", "ad_spend", "mer",
                "blended_cac", "meta_roas", "google_roas", "meta_cac", "google_cac", "refund_rate", "gross_margin_pct",
                "meta_roas_first6", "meta_roas_last6", "meta_cac_first6", "meta_cac_last6", "google_roas_last6",
                "google_cac_last6", "breakeven_roas", "worst_refund_sku", "worst_refund_rate", "m1_retention",
                "repeat_bedding_first", "repeat_other_first", "nov_discount_rate", "other_discount_rate"]:
        v = k[key]
        print(f"{key:24s} {v:,.4f}" if isinstance(v, float) else f"{key:24s} {v}")


if __name__ == "__main__":
    main()
